"""檢查簡報形狀是否交疊，或超出頁面、壓到右下角頁碼。

  python deck/check_layout.py deck/初賽簡報_vX.Y.pptx

包含關係（文字框在卡片裡、小標在大方塊裡）不算重疊。
v3.7 起沒有頁尾來源列：內容區到 7.36 吋，頁碼框在 12.92–13.30 × 7.16–7.40 吋（內容區右緣之外）。
「壓到頁碼」＝任何不是頁碼本身的形狀與頁碼框相交超過 0.04 吋。
卡片內留白：文字實際高度（依字級與換行估算）除以底下卡片高度，
低於 0.7 的列出來。高度不到 0.7 吋的小卡、頁尾與頁首標籤不列入。
名稱以 chart 開頭的形狀是原生圖表的圖區、長條與熱圖格，不是文字卡，也不列入。
若同版預覽圖存在，另外印出每張的 PIL 空白比例：
內容區（約 1.05 吋到 7.36 吋）裡 R、G、B 都 ≥ 250 的像素占比。
淺底色卡片不算白。
"""
from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU = 914400
OVERLAP_IN = 0.04
CONTAIN_TOL = 0.03
FOOTER_Y = 7.16  # 頁碼框上緣
BOT = 7.36  # 內容區下緣（build_deck.BOT）
PAGE_BOX = (12.92, FOOTER_Y, 13.30, FOOTER_Y + 0.24)  # 頁碼框（build_deck.chrome）
SLIDE_PAD = 0.02


def inches(value) -> float:
    return float(value) / EMU


def shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return " ".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())


def walk(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(shape.shapes)
            continue
        yield shape


def box_of(shape):
    return (
        inches(shape.left),
        inches(shape.top),
        inches(shape.left) + inches(shape.width),
        inches(shape.top) + inches(shape.height),
    )


def contains(outer, inner, tol=CONTAIN_TOL) -> bool:
    return (
        outer[0] <= inner[0] + tol
        and outer[1] <= inner[1] + tol
        and outer[2] >= inner[2] - tol
        and outer[3] >= inner[3] - tol
    )


def intersection(a, b):
    x0, y0 = max(a[0], b[0]), max(a[1], b[1])
    x1, y1 = min(a[2], b[2]), min(a[3], b[3])
    return max(0.0, x1 - x0), max(0.0, y1 - y0)


def label(shape, text) -> str:
    snippet = text.replace("\n", " ")[:28]
    kind = str(shape.shape_type).split(".")[-1]
    return f"{kind}:{snippet}" if snippet else kind


def check_slide(slide, index: int, slide_w: float, slide_h: float):
    items = []
    for shape in walk(slide.shapes):
        try:
            rect = box_of(shape)
        except (TypeError, AttributeError):
            continue
        items.append((shape, rect, shape_text(shape)))
    overlaps = []
    for i, (sa, ra, ta) in enumerate(items):
        for sb, rb, tb in items[i + 1 :]:
            iw, ih = intersection(ra, rb)
            if iw <= OVERLAP_IN or ih <= OVERLAP_IN:
                continue
            if contains(ra, rb) or contains(rb, ra):
                continue
            overlaps.append((label(sa, ta), label(sb, tb), iw, ih))
    outside = []
    for shape, rect, text in items:
        x0, y0, x1, y1 = rect
        reasons = []
        if x0 < -SLIDE_PAD or y0 < -SLIDE_PAD or x1 > slide_w + SLIDE_PAD or y1 > slide_h + SLIDE_PAD:
            reasons.append("超出頁面")
        # 頁碼本身（名稱 page-number）不算；其他形狀碰到頁碼框就列出。
        if (shape.name or "") != "page-number":
            iw, ih = intersection(rect, PAGE_BOX)
            if iw > OVERLAP_IN and ih > OVERLAP_IN:
                reasons.append("壓到頁碼")
        if reasons:
            outside.append((label(shape, text), ",".join(reasons), rect))
    return overlaps, outside


def _pt(run) -> float:
    size = run.font.size
    return float(size.pt) if size is not None else 14.0


def _char_width(ch: str, pt: float) -> float:
    o = ord(ch)
    if ch.isspace():
        return pt / 72.0 * 0.33
    if o > 0x2E00 or ch in "，。、；：？！（）「」『』％":
        return pt / 72.0
    return pt / 72.0 * 0.55


def glyph_height(shape) -> float:
    """依字級、邊界與換行估算文字高度（吋），不含文字框多餘的空白。"""
    if not getattr(shape, "has_text_frame", False):
        return 0.0
    tf = shape.text_frame
    width = inches(shape.width) - inches(tf.margin_left or 0) - inches(tf.margin_right or 0)
    if width < 0.25:
        return 0.0
    height = inches(tf.margin_top or 0) + inches(tf.margin_bottom or 0)
    for para in tf.paragraphs:
        text = "".join(run.text or "" for run in para.runs)
        if not text.strip():
            continue
        sizes = [_pt(run) for run in para.runs if (run.text or "").strip()]
        pt = max(sizes) if sizes else 14.0
        lines = 1
        used = 0.0
        for ch in text:
            w = _char_width(ch, pt)
            if used + w > width and used > 0:
                lines += 1
                used = w
            else:
                used += w
        height += lines * (pt / 72.0 * 1.15) + (2.0 / 72.0)
    return height


def card_fill(items):
    """一張卡片裡，上下堆疊的文字高度（並排取較高者）除以卡片高度。"""
    cards = []
    texts = []
    for shape, rect, text in items:
        h = rect[3] - rect[1]
        if (shape.name or "").startswith("chart"):
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and h >= 0.70 and rect[1] < BOT - 0.05:
            cards.append(rect)
            continue
        if not text.strip():
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX and h >= 0.30:
            texts.append((rect, glyph_height(shape), text))
    flagged = []
    frames = []
    for card in cards:
        if any(
            other is not card and contains(card, other, tol=0.02) and (other[3] - other[1]) >= 0.45
            for other in cards
        ):
            frames.append(card)
    for card in cards:
        if card in frames:
            continue
        ch = card[3] - card[1]
        inside = [(rect, gh, text) for rect, gh, text in texts if contains(card, rect, tol=0.08)]
        if not inside:
            continue
        inside.sort(key=lambda item: item[0][1])
        bands = []
        for rect, gh, text in inside:
            placed = False
            for band in bands:
                if rect[1] < band["bottom"] - 0.05 and rect[3] > band["top"] + 0.05:
                    band["gh"] = max(band["gh"], gh)
                    band["top"] = min(band["top"], rect[1])
                    band["bottom"] = max(band["bottom"], rect[3])
                    if len(text) > len(band["text"]):
                        band["text"] = text
                    placed = True
                    break
            if not placed:
                bands.append({"top": rect[1], "bottom": rect[3], "gh": gh, "text": text})
        used = sum(band["gh"] for band in bands)
        ratio = used / ch if ch else 1.0
        if ratio < 0.70:
            label = bands[0]["text"].replace("\n", " ")[:22]
            flagged.append((ratio, used, ch, label))
    return flagged


def whitespace(png: Path, slide_w: float, slide_h: float) -> float:
    im = Image.open(png).convert("RGB")
    px_per_in = im.width / slide_w
    x0 = int(0.30 * px_per_in)
    x1 = int((slide_w - 0.30) * px_per_in)
    y0 = int(1.05 * px_per_in)
    y1 = int(min(BOT, slide_h) * px_per_in)
    crop = im.crop((x0, y0, x1, y1))
    white = 0
    total = crop.width * crop.height
    pixels = list(crop.get_flattened_data()) if hasattr(crop, "get_flattened_data") else list(crop.getdata())
    for r, g, b in pixels:
        if r >= 250 and g >= 250 and b >= 250:
            white += 1
    return 100.0 * white / total if total else 0.0


def preview_dir(pptx: Path) -> Path | None:
    name = pptx.stem
    marker = "_v"
    if marker not in name:
        return None
    version = name.rsplit(marker, 1)[1]
    folder = pptx.parent / "preview" / f"v{version}"
    return folder if folder.is_dir() else None


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) != 2:
        print("用法：python deck/check_layout.py deck/初賽簡報_vX.Y.pptx")
        return 2
    path = Path(sys.argv[1])
    prs = Presentation(str(path))
    slide_w = inches(prs.slide_width)
    slide_h = inches(prs.slide_height)
    overlap_n = 0
    outside_n = 0
    card_n = 0
    for i, slide in enumerate(prs.slides, start=1):
        overlaps, outside = check_slide(slide, i, slide_w, slide_h)
        items = []
        for shape in walk(slide.shapes):
            try:
                items.append((shape, box_of(shape), shape_text(shape)))
            except (TypeError, AttributeError):
                continue
        sparse = card_fill(items)
        overlap_n += len(overlaps)
        outside_n += len(outside)
        card_n += len(sparse)
        if not overlaps and not outside and not sparse:
            continue
        print(f"第 {i} 頁")
        for a, b, iw, ih in overlaps:
            print(f"  重疊 {iw:.2f}x{ih:.2f} 吋  {a}  /  {b}")
        for name, reason, rect in outside:
            print(f"  {reason}  {name}  ({rect[0]:.2f},{rect[1]:.2f})-({rect[2]:.2f},{rect[3]:.2f})")
        for ratio, used, ch, label in sparse:
            print(f"  卡片內 {ratio:.2f}  文字 {used:.2f} / 卡片 {ch:.2f}  {label}")
    print(
        f"重疊 {overlap_n} 處，超出或壓線 {outside_n} 處，"
        f"卡片內低於 0.7 共 {card_n} 處，共 {len(prs.slides)} 頁"
    )
    folder = preview_dir(path)
    if folder is None:
        print("找不到預覽圖，略過 PIL 空白")
        return 0
    print("PIL 空白（內容區 RGB≥250）")
    for png in sorted(folder.glob("slide-*.png")):
        ratio = whitespace(png, slide_w, slide_h)
        print(f"  {png.name}  {ratio:.1f}%")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
