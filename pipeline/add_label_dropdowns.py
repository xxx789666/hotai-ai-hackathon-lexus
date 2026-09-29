"""替 human_labeling_A/B.xlsx 的 H（流失）與 I（面向）欄加下拉選單；保留已填內容。
I 欄清單為九面向＋無，多選仍可手打「價格;零件供應」（驗證不擋非清單值）。
  python pipeline/add_label_dropdowns.py
"""
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

P = Path(__file__).resolve().parent.parent / "data/processed"
ASPECTS = ["價格", "報價透明", "態度", "技術品質", "等待預約", "保固延保", "零件供應", "便利設施", "銷售交車", "無"]

for name in ("human_labeling_A.xlsx", "human_labeling_B.xlsx"):
    wb = load_workbook(P / name)
    for sheet in ("校準30", "正式300"):
        ws = wb[sheet]
        ws.data_validations.dataValidation = []  # 清掉舊的再加
        last = ws.max_row
        dv_h = DataValidation(type="list", formula1='"是,否"', allow_blank=True, showErrorMessage=True,
                              errorTitle="流失", error="請選 是 或 否")
        dv_i = DataValidation(type="list", formula1='"' + ",".join(ASPECTS) + '"', allow_blank=True,
                              showErrorMessage=False)  # 不擋多選手打
        dv_i.promptTitle = "面向"; dv_i.prompt = "單選用下拉；多選手打，用分號隔開，例：價格;零件供應"; dv_i.showInputMessage = True
        ws.add_data_validation(dv_h); dv_h.add(f"H2:H{last}")
        ws.add_data_validation(dv_i); dv_i.add(f"I2:I{last}")
    wb.save(P / name)
    print("ok", name)
