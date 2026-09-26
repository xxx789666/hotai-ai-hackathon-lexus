"""Sprint 1 PB-06：讀兩位標註者的 xlsx，算一致性，產生仲裁表，仲裁後算各模型對人工金標的指標。

  python pipeline/human_eval.py            # 讀 human_labeling_A.xlsx / _B.xlsx 的「正式300」，算 A/B κ，輸出仲裁表
  python pipeline/human_eval.py --final    # 讀 human_labels_arbitration.xlsx 的「最終」欄，算模型 vs 人工

輸入：data/processed/human_labeling_A.xlsx、human_labeling_B.xlsx、human_label_manifest.jsonl
輸出：data/processed/human_labels_arbitration.xlsx（不一致句在前，含「最終流失」「最終面向」欄）
      reports/T9_human_eval.md
"""
import argparse
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from verify_churn import load_jsonl  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
P = ROOT / "data/processed"
ASPECTS = ["價格", "報價透明", "態度", "技術品質", "等待預約", "保固延保", "零件供應", "便利設施", "銷售交車"]


def norm_churn(v):
    v = (str(v) if v is not None else "").strip()
    if v in ("是", "1", "Y", "y", "true", "True"):
        return 1
    if v in ("否", "0", "N", "n", "false", "False"):
        return 0
    return None


def norm_aspects(v):
    v = (str(v) if v is not None else "").strip().replace("，", ";").replace("、", ";").replace("；", ";")
    items = {x.strip() for x in v.split(";") if x.strip()}
    items = {x for x in items if x in ASPECTS}
    return items


def read_sheet(path, sheet="正式300"):
    ws = load_workbook(path, data_only=True)[sheet]
    rows = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[1]:
            continue
        sid = str(r[1])
        rows[sid] = {"churn": norm_churn(r[7]), "aspects": norm_aspects(r[8]),
                     "unsure": str(r[9] or "").strip() in ("1", "1.0"), "note": r[10] or "",
                     "row": list(r)}
    return rows


def kappa(pairs):
    """pairs: list of (a, b) binary."""
    n = len(pairs)
    if n == 0:
        return float("nan"), 0.0
    agree = sum(1 for a, b in pairs if a == b) / n
    pa = sum(a for a, _ in pairs) / n
    pb = sum(b for _, b in pairs) / n
    pe = pa * pb + (1 - pa) * (1 - pb)
    return (agree - pe) / (1 - pe) if pe < 1 else float("nan"), agree


def prf(pairs):
    """pairs: list of (gold, pred) binary."""
    tp = sum(1 for g, p in pairs if g and p)
    fp = sum(1 for g, p in pairs if not g and p)
    fn = sum(1 for g, p in pairs if g and not p)
    tn = sum(1 for g, p in pairs if not g and not p)
    prec = tp / (tp + fp) if tp + fp else 0.0
    rec = tp / (tp + fn) if tp + fn else 0.0
    f1 = 2 * prec * rec / (prec + rec) if prec + rec else 0.0
    k, acc = kappa([(p, g) for g, p in pairs])
    return dict(n=len(pairs), tp=tp, fp=fp, fn=fn, tn=tn, P=prec, R=rec, F1=f1, kappa=k, acc=acc)


def fmt(d):
    return f"| {d['n']} | {d['P']:.3f} | {d['R']:.3f} | {d['F1']:.3f} | {d['kappa']:.3f} | {d['tp']}/{d['fp']}/{d['fn']}/{d['tn']} |"


def stage1():
    A = read_sheet(P / "human_labeling_A.xlsx")
    B = read_sheet(P / "human_labeling_B.xlsx")
    sids = [s for s in A if s in B]
    manifest = {m["sid"]: m for m in load_jsonl(P / "human_label_manifest.jsonl")}
    pairs = [(A[s]["churn"], B[s]["churn"]) for s in sids if A[s]["churn"] is not None and B[s]["churn"] is not None]
    k, agree = kappa(pairs)
    lines = [f"# T9 人工標註一致性（第一階段）", "",
             f"共同標完 {len(pairs)} 句（A 未標 {sum(1 for s in sids if A[s]['churn'] is None)}、B 未標 {sum(1 for s in sids if B[s]['churn'] is None)}）。", "",
             f"流失二元：一致率 {agree:.3f}，Cohen's κ **{k:.3f}**。", ""]
    for stratum in ("A_random", "B_hard"):
        sub = [(A[s]["churn"], B[s]["churn"]) for s in sids if manifest[s]["stratum"] == stratum
               and A[s]["churn"] is not None and B[s]["churn"] is not None]
        k2, a2 = kappa(sub)
        lines.append(f"- {stratum}：n={len(sub)}，一致率 {a2:.3f}，κ {k2:.3f}")
    lines += ["", "## 面向逐項 κ", "", "| 面向 | A 標 | B 標 | 一致率 | κ |", "| --- | ---: | ---: | ---: | ---: |"]
    for asp in ASPECTS + ["無"]:
        sub = []
        for s in sids:
            a = (asp in A[s]["aspects"]) if asp != "無" else (not A[s]["aspects"])
            b = (asp in B[s]["aspects"]) if asp != "無" else (not B[s]["aspects"])
            sub.append((int(a), int(b)))
        k3, a3 = kappa(sub)
        lines.append(f"| {asp} | {sum(a for a, _ in sub)} | {sum(b for _, b in sub)} | {a3:.3f} | {k3:.3f} |")
    # 仲裁表
    wb = Workbook()
    ws = wb.active
    ws.title = "仲裁"
    ws.append(["sid", "來源", "標題", "前文", "目標句", "後文", "A流失", "B流失", "A面向", "B面向", "A備註", "B備註",
               "不一致", "最終流失(是/否)", "最終面向", "仲裁備註"])
    order = sorted(sids, key=lambda s: (0 if (A[s]["churn"] != B[s]["churn"] or A[s]["aspects"] != B[s]["aspects"]) else 1, s))
    n_dis = 0
    for s in order:
        r = A[s]["row"]
        dis = A[s]["churn"] != B[s]["churn"] or A[s]["aspects"] != B[s]["aspects"]
        n_dis += dis
        ws.append([s, r[2], r[3], r[4], r[5], r[6], r[7], B[s]["row"][7], ";".join(sorted(A[s]["aspects"])),
                   ";".join(sorted(B[s]["aspects"])), A[s]["note"], B[s]["note"], "1" if dis else "",
                   "" if dis else r[7], "" if dis else ";".join(sorted(A[s]["aspects"])), ""])
    wb.save(P / "human_labels_arbitration.xlsx")
    lines += ["", f"不一致（流失或面向任一不同）{n_dis} 句，已排在 `human_labels_arbitration.xlsx` 最前面；一致者的最終欄已預填。仲裁完成後執行 `--final`。"]
    (ROOT / "reports/T9_human_eval.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("\n".join(lines))


def stage2():
    ws = load_workbook(P / "human_labels_arbitration.xlsx", data_only=True)["仲裁"]
    final = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        c = norm_churn(r[13])
        if c is None:
            continue
        final[str(r[0])] = {"churn": c, "aspects": norm_aspects(r[14])}
    manifest = {m["sid"]: m for m in load_jsonl(P / "human_label_manifest.jsonl")}
    r4 = {}
    for name in ("jev_lora_r4_ep2_pilot.jsonl", "jev_lora_r4_ep2_val.jsonl"):
        for r in load_jsonl(P / name):
            p = r["churn_p"]
            r4[r["sid"]] = (p[2] + p[3]) if isinstance(p, list) else (p.get("2", 0) + p.get("3", 0))
    haiku_asp = {r["sid"]: set(r.get("a", [])) for r in load_jsonl(P / "aspect_pilot_haiku.jsonl")} if (P / "aspect_pilot_haiku.jsonl").exists() else {}
    models = {
        "Sonnet／GPT-5.6 Sol（金標來源）": lambda s: int(manifest[s]["gold_c"] >= 2),
        "Haiku 零樣本": lambda s: int(manifest[s]["haiku_c"] >= 2),
        "r4 c_expect": lambda s: int((manifest[s]["r4_c"] or 0) >= 2),
        "r4 P(c≥2)≥0.5": lambda s: int(r4.get(s, 0) >= 0.5),
        "r4 P(c≥2)≥0.3": lambda s: int(r4.get(s, 0) >= 0.3),
    }
    lines = ["# T9 人工金標 vs 各模型（第二階段）", "", f"仲裁完成 {len(final)} 句。人工金標＝仲裁後的「最終流失」。", ""]
    for stratum, title in (("A_random", "隨機層（母體估計，簡報引用）"), ("B_hard", "困難層（模型盲點）"), (None, "全部 300 句")):
        sids = [s for s in final if stratum is None or manifest[s]["stratum"] == stratum]
        pos = sum(final[s]["churn"] for s in sids)
        lines += [f"## {title}：n={len(sids)}，人工判流失 {pos}（{pos / max(len(sids), 1):.1%}）", "",
                  "| 模型 | n | P | R | F1 | κ | TP/FP/FN/TN |", "| --- | ---: | ---: | ---: | ---: | ---: | --- |"]
        for name, fn in models.items():
            d = prf([(final[s]["churn"], fn(s)) for s in sids])
            lines.append(f"| {name} " + fmt(d))
        lines.append("")
    # 面向：Sonnet vs 人工
    lines += ["## 面向：Sonnet 標註 vs 人工（全部句）", "", "| 面向 | 人工有 | Sonnet 有 | P | R | F1 |", "| --- | ---: | ---: | ---: | ---: | ---: |"]
    micro = Counter()
    for asp in ASPECTS:
        pairs = [(int(asp in final[s]["aspects"]), int(asp in set(manifest[s]["gold_a"]))) for s in final]
        d = prf(pairs)
        micro["tp"] += d["tp"]; micro["fp"] += d["fp"]; micro["fn"] += d["fn"]
        lines.append(f"| {asp} | {sum(g for g, _ in pairs)} | {sum(p for _, p in pairs)} | {d['P']:.2f} | {d['R']:.2f} | {d['F1']:.2f} |")
    mp = micro["tp"] / max(micro["tp"] + micro["fp"], 1); mr = micro["tp"] / max(micro["tp"] + micro["fn"], 1)
    lines.append(f"\nSonnet 面向 micro F1 = {2 * mp * mr / max(mp + mr, 1e-9):.3f}（P {mp:.2f}，R {mr:.2f}）")
    if haiku_asp:
        micro = Counter()
        sids_h = [s for s in final if s in haiku_asp]
        for asp in ASPECTS:
            pairs = [(int(asp in final[s]["aspects"]), int(asp in haiku_asp[s])) for s in sids_h]
            d = prf(pairs); micro["tp"] += d["tp"]; micro["fp"] += d["fp"]; micro["fn"] += d["fn"]
        mp = micro["tp"] / max(micro["tp"] + micro["fp"], 1); mr = micro["tp"] / max(micro["tp"] + micro["fn"], 1)
        lines.append(f"Haiku 面向 micro F1 = {2 * mp * mr / max(mp + mr, 1e-9):.3f}（n={len(sids_h)} 句有 Haiku 面向）")
    (ROOT / "reports/T9_human_eval.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("\n".join(lines))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--final", action="store_true")
    args = ap.parse_args()
    stage2() if args.final else stage1()
