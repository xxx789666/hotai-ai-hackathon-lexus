"""Sprint 1 PB-06：產生人工標註材料（校準 30 句 + 正式 300 句）。

候選池 = r4 未參與訓練的句子（600 句測試集 + 1,096 句驗證集，共 1,696 句）。
  隨機層 A：150 句，seed 42 簡單隨機，用來估母體指標。
  困難層 B：150 句，從「r4 機率 0.2–0.8」「r4 與金標不一致」「Haiku 與金標不一致」「T3 疑似誤標」的聯集抽，
           優先順序：疑似誤標 > r4 不一致 > Haiku 不一致 > 不確定區。
  校準 C：30 句，池中剩餘句抽 15 金標流失 + 15 非流失（其中 8 句取自困難句）。
文字一律經 deidentify.mask_text 遮蔽；金標與模型預測不寫進 xlsx，只留在 manifest。

  python pipeline/build_human_labeling.py
輸出：
  data/processed/human_label_manifest.jsonl   sid、層別、金標與各模型預測（評估用，不給標註者）
  data/processed/human_labeling_A.xlsx / _B.xlsx   兩位標註者各一份（內容相同、順序相同）
"""
import json
import random
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deidentify import compile_shop_map, load_salt, mask_text, scan_shops  # noqa: E402
from verify_churn import build_context, load_jsonl  # noqa: E402

P = Path(__file__).resolve().parent.parent / "data/processed"
ASPECTS = ["價格", "報價透明", "態度", "技術品質", "等待預約", "保固延保", "零件供應", "便利設施", "銷售交車"]
SEED = 42


def p2(row):
    p = row["churn_p"]
    return (p[2] + p[3]) if isinstance(p, list) else (p.get("2", 0) + p.get("3", 0))


def main():
    rnd = random.Random(SEED)
    gold = {}
    for r in load_jsonl(P / "churn_verified.jsonl"):
        gold[r["sid"]] = r
    haiku = {r["sid"]: r["c"] for r in load_jsonl(P / "churn_labels.jsonl")}
    aspects = {r["sid"]: r for r in load_jsonl(P / "aspect_labels.jsonl")}
    t2 = {r["sid"]: r.get("t2") for r in load_jsonl(P / "churn_positives.jsonl")}
    sents = {r["sid"]: r for r in load_jsonl(P / "aftersales.jsonl")}

    pool = {}
    for name, tag in (("jev_lora_r4_ep2_pilot.jsonl", "test"), ("jev_lora_r4_ep2_val.jsonl", "val")):
        for r in load_jsonl(P / name):
            pool[r["sid"]] = {"set": tag, "r4_p": round(p2(r), 4), "r4_c": r.get("c_expect", r.get("c"))}
    sids = sorted(pool)
    print(f"pool {len(sids)}")

    def is_pos(s):
        return gold[s]["c"] >= 2

    sus = [s for s in sids if t2.get(s) == "疑似誤標"]
    dis_r4 = [s for s in sids if (pool[s]["r4_p"] >= 0.5) != is_pos(s)]
    dis_h = [s for s in sids if (haiku[s] >= 2) != is_pos(s)]
    unc = [s for s in sids if 0.2 <= pool[s]["r4_p"] < 0.8]
    hard_reason = {}
    for s in sus:
        hard_reason.setdefault(s, "疑似誤標")
    for s in dis_r4:
        hard_reason.setdefault(s, "r4不一致")
    for s in dis_h:
        hard_reason.setdefault(s, "haiku不一致")
    for s in unc:
        hard_reason.setdefault(s, "r4不確定")

    # A 隨機層
    A = rnd.sample(sids, 150)
    taken = set(A)
    # B 困難層：依優先順序，每級內隨機
    B = []
    for reason in ("疑似誤標", "r4不一致", "haiku不一致", "r4不確定"):
        cand = [s for s in sids if hard_reason.get(s) == reason and s not in taken]
        rnd.shuffle(cand)
        for s in cand:
            if len(B) >= 150:
                break
            B.append(s)
            taken.add(s)
    # C 校準 30：15 正 15 負，負例中 8 句是困難句
    rest = [s for s in sids if s not in taken]
    pos_rest = [s for s in rest if is_pos(s)]
    neg_hard = [s for s in rest if not is_pos(s) and s in hard_reason]
    neg_easy = [s for s in rest if not is_pos(s) and s not in hard_reason]
    C = rnd.sample(pos_rest, 15) + rnd.sample(neg_hard, 8) + rnd.sample(neg_easy, 7)
    rnd.shuffle(C)
    main300 = A + B
    rnd.shuffle(main300)

    # 去識別
    salt = load_salt()
    counts = Counter()
    _rows, shop_pat = compile_shop_map(scan_shops(list(sents.values())))
    ctx = build_context(main300 + C)

    def mask(t):
        return mask_text(t or "", counts, shop_pat)

    manifest = []
    for stratum, group in (("A_random", A), ("B_hard", B), ("C_calib", C)):
        for s in group:
            g = gold[s]
            manifest.append({
                "sid": s, "stratum": stratum, "set": pool[s]["set"],
                "hard_reason": hard_reason.get(s), "gold_c": g["c"], "gold_w": g.get("w"),
                "gold_model": g.get("model"), "gold_a": aspects.get(s, {}).get("a", []),
                "haiku_c": haiku[s], "r4_p": pool[s]["r4_p"], "r4_c": pool[s]["r4_c"],
                "t2": t2.get(s),
            })
    with open(P / "human_label_manifest.jsonl", "w", encoding="utf-8") as f:
        for m in manifest:
            f.write(json.dumps(m, ensure_ascii=False) + "\n")

    def build_xlsx(path, labeler):
        wb = Workbook()
        ws0 = wb.active
        ws0.title = "說明"
        lines = [
            f"標註者：{labeler}", "",
            "每句只標兩欄：",
            "  流失：是／否。「是」＝目標句透露車主考慮離開或已經離開 Lexus 原廠保養維修（含建議別人離開、因售後想換掉 Lexus）。只抱怨沒行動、單純詢問或閒聊都是「否」。",
            "  面向：目標句談到的售後面向，可多選，用分號隔開；沒談到售後就填「無」。",
            "  選項：價格；報價透明；態度；技術品質；等待預約；保固延保；零件供應；便利設施；銷售交車；無",
            "只看「目標句」，前後文與標題只用來理解語境。",
            "拿不準就在「不確定」填 1，備註寫一句為什麼；不要空著。",
            "先做「校準30」，兩人對答案、討論分歧後，再各自獨立做「正式300」，做的時候不要討論。",
            "完整定義與邊界例見 標註指引_2026-09-29.md。",
        ]
        for i, t in enumerate(lines, 1):
            ws0.cell(row=i, column=1, value=t)
        ws0.column_dimensions["A"].width = 120

        def fill(ws, group):
            head = ["序號", "sid", "來源", "標題", "前文", "目標句", "後文", "流失(是/否)", "面向(分號隔開)", "不確定(1)", "備註"]
            ws.append(head)
            for c in range(1, len(head) + 1):
                ws.cell(row=1, column=c).font = Font(bold=True)
                ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor="DDDDDD")
            for i, s in enumerate(group, 1):
                c = ctx.get(s, {"title": sents[s]["title"], "prev": [], "text": sents[s]["text"], "next": None})
                ws.append([i, s, sents[s]["source"], mask(c["title"]), mask(" / ".join(c["prev"])),
                           mask(c["text"]), mask(c["next"] or ""), "", "", "", ""])
            widths = {"A": 6, "B": 16, "C": 9, "D": 28, "E": 40, "F": 50, "G": 40, "H": 12, "I": 26, "J": 10, "K": 30}
            for k, v in widths.items():
                ws.column_dimensions[k].width = v
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            dv = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"H2:H{len(group) + 1}")
            ws.freeze_panes = "F2"

        ws1 = wb.create_sheet("校準30")
        fill(ws1, C)
        ws2 = wb.create_sheet("正式300")
        fill(ws2, main300)
        wb.save(path)

    build_xlsx(P / "human_labeling_A.xlsx", "A")
    build_xlsx(P / "human_labeling_B.xlsx", "B")
    print("A", len(A), "B", len(B), "C", len(C))
    print("B reasons", Counter(hard_reason[s] for s in B))
    print("gold pos: A", sum(is_pos(s) for s in A), "B", sum(is_pos(s) for s in B), "C", sum(is_pos(s) for s in C))
    print("mask counts", dict(counts))


if __name__ == "__main__":
    main()
